"""Выгрузка общей базы «Уязвимостей» в Excel — в стиле AuditLens.

Книга как у «Отзывов»: «Обзор» (фирменная шапка, штамп «Экспортировано из
AuditLens», фильтры, плитки и родные графики Excel), «Записи» (всё, что видно
в базе по фильтру или отмечено: даты — датами, вероятность — числом, вердикт и
проверка подсвечены цветами интерфейса) и «Сводки» — таблицы под графиками.
Выгружаются записи с любым статусом: на проде все записи «предварительные»,
и выгрузка только опубликованных была пустой.
"""
from __future__ import annotations

import io
from collections import Counter
from datetime import date, datetime

from ..web.export_brand import MONO, SERIF, XL, C, banner_png, stamp_line

BANK_NAMES = {
    "sberbank": "Сбербанк", "sber": "Сбербанк", "vtb": "ВТБ", "alfabank": "Альфа-Банк",
    "alfa": "Альфа-Банк", "tbank": "Т-Банк", "tinkoff": "Т-Банк", "gazprombank": "Газпромбанк",
    "gpb": "Газпромбанк", "raiffeisen": "Райффайзенбанк", "rosbank": "Росбанк",
    "sovcombank": "Совкомбанк", "mtsbank": "МТС Банк", "mts": "МТС Банк",
    "pochtabank": "Почта Банк", "otkritie": "Открытие", "psb": "ПСБ",
    "rshb": "Россельхозбанк", "domrf": "Банк ДОМ.РФ", "ozon": "Озон Банк",
    "ozonbank": "Озон Банк", "yandex": "Яндекс Банк", "uralsib": "Уралсиб",
    "akbars": "Ак Барс", "mkb": "МКБ", "homecredit": "Хоум Банк",
    "renaissance": "Ренессанс Банк", "all": "Все банки", "generic": "Банк не указан",
    "other": "Другие банки",
}
TYPE_LABELS = {
    "vulnerability": "Уязвимость", "fraud_scheme": "Мошенническая схема",
    "not_confirmed": "Не подтверждено", None: "Без вердикта",
}
TYPE_COLORS = {"vulnerability": "accent", "fraud_scheme": "legal", "not_confirmed": "ink3"}
VERIFICATION_LABELS = {"awaiting": "Ждут проверки", "reviewed": "Проверено"}
CLASSIFICATION_FILTERS = {
    "confirmed": "Уязвимости и мошеннические схемы", "vulnerability": "Уязвимости",
    "fraud_scheme": "Мошеннические схемы", "not_confirmed": "Не подтверждено",
    "all": "Все записи",
}

COLUMNS = [
    ("published", "Дата публикации", 13), ("collected", "Собрано", 13),
    ("bank", "Банк", 16), ("type", "Вердикт", 20), ("check", "Проверка", 16),
    ("confidence", "Вероятность", 12), ("title", "Запись", 48),
    ("summary", "Суть / комментарий классификатора", 60),
    ("doubt", "Сомнение модели", 28), ("domain", "Источник", 18),
    ("url", "Ссылка", 40),
]


def bank_name(slug: str | None) -> str:
    value = str(slug or "").strip()
    return BANK_NAMES.get(value.lower(), value) if value else "Банк не указан"


def _kind(record: dict) -> str | None:
    kind = record.get("classification")
    if kind:
        return kind
    if record.get("is_loophole") is True:
        return "vulnerability"
    if record.get("is_loophole") is False:
        return "not_confirmed"
    return None


def _dt(value):
    if not value:
        return None
    if isinstance(value, datetime):
        return value.replace(tzinfo=None)
    if isinstance(value, date):
        return datetime(value.year, value.month, value.day)
    try:
        return datetime.fromisoformat(str(value).replace("Z", "+00:00")).replace(tzinfo=None)
    except ValueError:
        return None


def _check(record: dict) -> str:
    if record.get("reviewed"):
        return "Проверено"
    if record.get("awaiting"):
        return "Ждёт проверки"
    return "—"


def describe_filters(filters: dict) -> list[tuple[str, str]]:
    banks = filters.get("bank_slugs") or []
    period = "всё время"
    if filters.get("period_from") or filters.get("period_to"):
        period = f"{filters.get('period_from') or '…'} — {filters.get('period_to') or '…'}"
    return [
        ("Тип записи", CLASSIFICATION_FILTERS.get(filters.get("classification") or "all",
                                                   "Все записи")),
        ("Банки", ", ".join(bank_name(b) for b in banks) if banks else "все банки"),
        ("Дата публикации", period),
        ("Проверка", VERIFICATION_LABELS.get(filters.get("verification_status"), "все")),
        ("Поиск", f"«{filters['q']}»" if filters.get("q") else "без поиска"),
        ("Отбор", f"отмечено вручную: {len(filters['record_ids'])}"
         if filters.get("record_ids") else "все записи по фильтру"),
    ]


def to_xlsx(records: list[dict], filters: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    from openpyxl.chart.label import DataLabelList
    from openpyxl.styles import Alignment
    from openpyxl.worksheet.pagebreak import Break

    kinds = Counter(_kind(r) for r in records)
    banks = Counter(bank_name(r.get("bank_slug")) for r in records)
    awaiting = sum(1 for r in records if r.get("awaiting"))
    reviewed = sum(1 for r in records if r.get("reviewed"))
    title = "Уязвимости и мошеннические схемы"
    fl = describe_filters(filters)
    sub = " · ".join([f"{len(records):,} записей".replace(",", " "), fl[0][1],
                      "вердикты модели и экспертов ЦК КС"])

    wb = Workbook()
    wb.properties.creator = "AuditLens"
    wb.properties.lastModifiedBy = "AuditLens"
    wb.properties.title = title
    wb.properties.subject = "Общая база — вкладка «Уязвимости»"
    wb.properties.description = stamp_line()
    wb.properties.keywords = "AuditLens"

    # ── «Сводки» ─────────────────────────────────────────────────────────
    ss = wb.active
    ss.title = "Сводки"
    XL.sheet_base(ss, widths=[30, 12, 3, 26, 12], title=title)
    head = ss.cell(row=1, column=1, value="Сводки по выгрузке")
    head.font = XL.font(18, name=SERIF)
    XL.stamp(ss, 2, span=5)
    type_rows = [[TYPE_LABELS.get(k, "Без вердикта"), n] for k, n in kinds.most_common()]
    bank_rows = [[b, n] for b, n in banks.most_common(15)]
    for col, name, rows_ in ((1, "По вердиктам", type_rows), (4, "По банкам", bank_rows)):
        cap = ss.cell(row=4, column=col, value=name.upper())
        cap.font = XL.font(8.5, color="accent", name=MONO)
        XL.table(ss, 5, ["Вердикт" if col == 1 else "Банк", "Записей"], rows_, col=col)
    ss.auto_filter.ref = None

    # ── «Обзор» ──────────────────────────────────────────────────────────
    ov = wb.create_sheet("Обзор", 0)
    XL.sheet_base(ov, widths=[11.5] * 12, title=title)
    row = XL.banner(ov, banner_png(eyebrow="Уязвимости · общая база", title=title,
                                   subtitle=sub, width=1180), rows=8, width_px=1060)
    XL.stamp(ov, row, text_=stamp_line() + " · вкладка «Уязвимости»", span=12)
    row += 1
    fc = ov.cell(row=row, column=1,
                 value="Фильтры: " + " · ".join(f"{k.lower()}: {v}" for k, v in fl))
    fc.font = XL.font(9, color="ink3")
    fc.alignment = Alignment(wrap_text=True, vertical="top")
    ov.merge_cells(start_row=row, start_column=1, end_row=row, end_column=12)
    ov.row_dimensions[row].height = 28
    row += 2
    row = XL.kpis(ov, row, [
        ("Записей", len(records), "в выгрузке"),
        ("Уязвимости", kinds.get("vulnerability", 0), "вердикт модели или эксперта"),
        ("Мошеннические схемы", kinds.get("fraud_scheme", 0), "вердикт модели или эксперта"),
        ("Не подтверждено", kinds.get("not_confirmed", 0), "признаков лазейки нет"),
        ("Ждут проверки", awaiting, "решения эксперта ЦК КС ещё нет"),
        ("Проверено", reviewed, "есть решение эксперта"),
    ])
    ov.row_breaks.append(Break(id=row))
    row = XL.section(ov, row + 1, "Что в выгрузке",
                     eyebrow="Графики · по данным листа «Сводки»", span=12)
    for idx, (col, name, n) in enumerate(((1, "По вердиктам", len(type_rows)),
                                          (4, "По банкам", len(bank_rows)))):
        if not n:
            continue
        ch = BarChart()
        ch.type = "bar"
        ch.add_data(Reference(ss, min_col=col + 1, min_row=5, max_row=5 + n),
                    titles_from_data=True)
        ch.set_categories(Reference(ss, min_col=col, min_row=6, max_row=5 + n))
        ch.gapWidth = 60
        ch.width, ch.height = 13.2, 8.4
        ch.x_axis.scaling.orientation = "maxMin"
        ch.y_axis.scaling.min = 0
        ch.dataLabels = DataLabelList(showVal=True, showSerName=False, showCatName=False,
                                      showLegendKey=False, showPercent=False)
        XL.style_chart(ch, title=name, colors=[C["accent"] if idx == 0 else C["sber"]],
                       horizontal=True)
        ov.add_chart(ch, f"{'A' if idx == 0 else 'G'}{row}")
    row += 18
    note = ov.cell(row=row, column=1, value=(
        "Вердикт — модель при сборе записи или эксперт ЦК КС; «Ждут проверки» — находки "
        "модели без решения эксперта. Вероятность — оценка модели, а не решение эксперта."))
    note.font = XL.font(8.5, color="ink3")
    note.alignment = Alignment(wrap_text=True, vertical="top")
    ov.merge_cells(start_row=row, start_column=1, end_row=row, end_column=12)
    ov.row_dimensions[row].height = 30

    # ── «Записи» ─────────────────────────────────────────────────────────
    ws = wb.create_sheet("Записи", 1)
    XL.sheet_base(ws, widths=[w for _, _, w in COLUMNS], title=title)
    t = ws.cell(row=1, column=1, value=title)
    t.font = XL.font(16, name=SERIF)
    XL.stamp(ws, 2, span=len(COLUMNS))
    data = []
    for r in records:
        kind = _kind(r)
        summary = r.get("summary") or r.get("classifier_verdict_reason") or r.get("verdict_reason")
        conf = r.get("verdict_confidence")
        data.append([
            _dt(r.get("published_at")), _dt(r.get("collected_at")),
            bank_name(r.get("bank_slug")), TYPE_LABELS.get(kind, "Без вердикта"), _check(r),
            float(conf) if conf is not None else None,
            r.get("headline") or r.get("title") or r.get("snippet") or "",
            summary or "", r.get("summary_doubt") or "", r.get("domain") or "", r.get("url") or "",
        ])
    XL.table(ws, 4, [label for _, label, _ in COLUMNS], data,
             formats={"Дата публикации": "dd.mm.yyyy hh:mm", "Собрано": "dd.mm.yyyy hh:mm",
                      "Вероятность": "0%"},
             wrap={"Запись", "Суть / комментарий классификатора", "Сомнение модели"},
             links={"Ссылка"})
    for i, r in enumerate(records):
        kind = _kind(r)
        cell = ws.cell(row=5 + i, column=4)
        cell.font = XL.font(10, color=TYPE_COLORS.get(kind, "ink3"), bold=kind in TYPE_COLORS
                            and kind != "not_confirmed")
        check = ws.cell(row=5 + i, column=5)
        if r.get("awaiting") and not r.get("reviewed"):
            check.font = XL.font(10, color="warn")
    ws.freeze_panes = "A5"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
