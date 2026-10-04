"""Экспорт результата loophole в PDF через Playwright.

Переиспользует стиль web/pdf_export (Source Serif 4 / Geist / JetBrains Mono).
Генерирует HTML из записей, рендерит в A4 PDF через headless Chromium.
"""
from __future__ import annotations

import logging
import re
from html import escape
from io import BytesIO

from .direct_transport import chromium_args

log = logging.getLogger(__name__)

# Поля страницы A4 по ГОСТ-подобной раскладке: слева 30 мм, справа 10 мм,
# сверху и снизу по 20 мм (нижнее поле — отступ от нижней грани листа).
_PDF_MARGINS = {"top": "20mm", "bottom": "20mm", "left": "30mm", "right": "10mm"}


_HTML_TEMPLATE = """<!DOCTYPE html>
<html lang="ru"><head><meta charset="utf-8">
<link rel="preconnect" href="https://fonts.googleapis.com">
<link rel="stylesheet" href="https://fonts.googleapis.com/css2?family=Source+Serif+4:opsz,wght@8..60,400;8..60,600&family=Geist:wght@400;500;600&family=JetBrains+Mono:wght@400;500&display=swap">
<style>
  body {{ font-family: "Geist", system-ui, sans-serif; color: #1a1a1a; margin: 0; padding: 0; }}
  h1 {{ font-family: "Source Serif 4", Georgia, serif; font-size: 1.5rem; break-after: avoid; }}
  h2 {{ font-family: "Source Serif 4", Georgia, serif; font-size: 1.1rem; margin-top: 18px; break-after: avoid; }}
  .meta {{ color: #6b6b6b; font-size: 0.85rem; margin-bottom: 20px; }}
  .record {{ border-bottom: 1px solid #d8d8d2; padding: 12px 0; break-inside: avoid; }}
  .record .title {{ font-weight: 600; }}
  .record .url {{ font-family: "JetBrains Mono", monospace; font-size: 0.8rem; color: #6b6b6b; word-break: break-all; }}
  .record .verdict {{ margin-top: 4px; }}
  .loophole {{ color: #b03a2e; }}
</style></head><body>
<h1>Отчёт: лазейки и мошеннические схемы</h1>
<div class="meta">Сформирован: {generated_at}</div>
{records_html}
</body></html>"""


def _record_html(r: dict) -> str:
    title = r.get("title") or r.get("snippet") or "(без названия)"
    url = r.get("url") or ""
    bank = r.get("bank_slug") or "—"
    is_l = r.get("is_loophole")
    verdict_cls = "loophole" if is_l else ""
    verdict = "лазейка" if is_l else "не лазейка"
    conf = r.get("verdict_confidence")
    conf_str = f" (доверие {conf:.2f})" if conf is not None else ""
    return (
        f'<div class="record">'
        f'<div class="title">{title}</div>'
        f'<div class="url">{url}</div>'
        f'<div>Банк: {bank}</div>'
        f'<div class="verdict {verdict_cls}">Вердикт: {verdict}{conf_str}</div>'
        f'</div>'
    )


def render_html(records: list[dict], *, generated_at: str = "") -> str:
    """Возвращает HTML отчёта."""
    from ..clock import today_ru
    gen = generated_at or today_ru()
    records_html = "\n".join(_record_html(r) for r in records)
    return _HTML_TEMPLATE.format(generated_at=gen, records_html=records_html)


async def export_pdf(records: list[dict], *, output_path: str = "") -> bytes:
    """Рендерит HTML в PDF через Playwright. Возвращает PDF-байты.

    Если Playwright недоступен — падает с понятной ошибкой.
    """
    html = render_html(records)
    try:
        from playwright.async_api import async_playwright
    except Exception as e:
        raise RuntimeError(f"Playwright недоступен: {e}") from e
    async with async_playwright() as p:
        browser = await p.chromium.launch(args=chromium_args())
        page = await browser.new_page()
        await page.set_content(html, wait_until="networkidle")
        pdf = await page.pdf(format="A4", print_background=True, margin=_PDF_MARGINS)
        await browser.close()
    if output_path:
        from pathlib import Path
        Path(output_path).write_bytes(pdf)
    return pdf


_RE_MD_LINK = re.compile(r"\[([^\]]+)\]\((https?://[^)\s]+)\)")
_RE_RAW_URL = re.compile(r"https?://[^\s<>()\[\]\"']+")


def render_research_report_html(report: dict) -> str:
    """Собирает отдельный безопасный HTML для immutable отчёта исследования.

    URL из текста итога убираются: вместо них подставляются номера [N]
    из нумерованного списка используемых источников в конце отчёта.
    Тексты статей (extracted_text) в отчёт не попадают.
    """
    from .markdown_render import render_markdown_html

    sources: list[dict[str, str]] = []
    source_index: dict[str, int] = {}

    def source_no(url: str, title: str = "") -> int:
        if url not in source_index:
            source_index[url] = len(sources) + 1
            sources.append({"url": url, "title": title})
        elif title and not sources[source_index[url] - 1]["title"]:
            sources[source_index[url] - 1]["title"] = title
        return source_index[url]

    def replace_link(match: re.Match[str]) -> str:
        label, url = match.group(1), match.group(2)
        return f"{label} [{source_no(url, label)}]"

    def replace_url(match: re.Match[str]) -> str:
        url = match.group(0).rstrip(".,;:!?")
        tail = match.group(0)[len(url):]
        return f"[{source_no(url)}]{tail}"

    # Итог исследования — markdown от аналитика: ссылки выносим в список
    # источников, затем рендерим разметку (вход экранируется внутри рендерера).
    result_text = _RE_MD_LINK.sub(replace_link, str(report.get("result") or ""))
    result_text = _RE_RAW_URL.sub(replace_url, result_text)
    result_html = render_markdown_html(result_text) or "<p>—</p>"

    for item in report.get("evidence") or []:
        if isinstance(item, dict) and item.get("url"):
            source_no(str(item["url"]), str(item.get("title") or ""))

    if sources:
        items = "".join(
            "<li>{title} — <a href=\"{url}\">ссылка</a></li>".format(
                title=escape(s["title"] or "Источник без названия"),
                url=escape(s["url"], quote=True),
            )
            for s in sources
        )
        sources_html = f"<ol>{items}</ol>"
    else:
        sources_html = "<p>Источники не использовались.</p>"

    def paragraph(value: object) -> str:
        return "<br>".join(escape(line) for line in str(value or "").splitlines()) or "—"

    return """<!doctype html><html lang=\"ru\"><head><meta charset=\"utf-8\"><style>
body {{ font-family: Arial, sans-serif; color: #1a1a1a; margin: 0; padding: 0; line-height: 1.55; }}
h1 {{ font-size: 24px; }} h2 {{ margin-top: 24px; }}
h1, h2, h3, h4, h5, h6 {{ break-after: avoid; }}
li, pre, blockquote, tr {{ break-inside: avoid; }}
h3 {{ font-size: 17px; margin: 18px 0 8px; }} h4, h5, h6 {{ font-size: 15px; margin: 14px 0 6px; }}
p {{ margin: 8px 0; }} ul, ol {{ margin: 8px 0; padding-left: 22px; }} li {{ margin: 3px 0; }}
a {{ color: #1a5fb4; }}
code {{ font-family: "JetBrains Mono", monospace; font-size: 0.86em; background: #f1f1ec; padding: 1px 4px; border-radius: 3px; }}
pre {{ background: #f1f1ec; border: 1px solid #d8d8d2; border-radius: 6px; padding: 10px 12px; overflow-x: auto; }}
pre code {{ background: none; padding: 0; }}
blockquote {{ margin: 10px 0; padding: 4px 14px; border-left: 3px solid #d8d8d2; color: #555; }}
hr {{ border: none; border-top: 1px solid #d8d8d2; margin: 16px 0; }}
.md-table-wrap {{ margin: 10px 0; }}
table {{ border-collapse: collapse; width: 100%; table-layout: fixed; font-size: 0.85em; }}
th, td {{ border: 1px solid #d8d8d2; padding: 6px 10px; text-align: left; vertical-align: top; word-break: break-word; overflow-wrap: anywhere; }}
th {{ background: #f6f6f2; }}
</style></head><body><h1>Отчёт AI-исследования</h1><h2>Тема</h2><p>{query}</p>
<h2>Итог</h2><div class="md-result">{result}</div><h2>Список используемых источников</h2>{sources}
</body></html>""".format(
        query=paragraph(report.get("query")), result=result_html, sources=sources_html
    )


async def export_research_report_pdf(report: dict) -> bytes:
    """Рендерит PDF отчёта исследования, не используя небезопасный catalog renderer."""
    try:
        from playwright.async_api import async_playwright
    except Exception as exc:
        raise RuntimeError(f"Playwright недоступен: {exc}") from exc
    async with async_playwright() as playwright:
        browser = await playwright.chromium.launch(args=chromium_args())
        try:
            page = await browser.new_page()
            await page.set_content(render_research_report_html(report), wait_until="networkidle")
            return await page.pdf(format="A4", print_background=True, margin=_PDF_MARGINS)
        finally:
            await browser.close()


def export_research_report_docx(report: dict) -> bytes:
    """Создаёт Word-документ из тех же immutable данных, что и PDF."""
    try:
        from docx import Document
    except Exception as exc:
        raise RuntimeError(f"Word-экспорт недоступен: {exc}") from exc
    document = Document()
    document.add_heading("Отчёт AI-исследования", level=0)
    document.add_heading("Тема", level=1)
    document.add_paragraph(str(report.get("query") or ""))
    document.add_heading("Итог", level=1)
    document.add_paragraph(str(report.get("result") or ""))
    document.add_heading("Проверенные доказательства и источники", level=1)
    evidence = report.get("evidence") or []
    if not evidence:
        document.add_paragraph("Проверенные доказательства отсутствуют.")
    for item in evidence:
        if not isinstance(item, dict):
            continue
        document.add_paragraph(str(item.get("title") or "Источник без названия"), style="List Bullet")
        document.add_paragraph(str(item.get("url") or ""))
        if item.get("extracted_text"):
            document.add_paragraph(str(item["extracted_text"]))
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()
