"""Фирменный стиль AuditLens для выгрузок в Excel и Word.

ЗАЧЕМ. Выгрузки жили своей жизнью: CSV без оформления, Excel с серой шапкой,
Word на Calibri — файл, отданный аудитором дальше, никак не выдавал, откуда
он и что это тот же инструмент. Владелец 26.09: «стиль инструмента должен
сохраняться в этих файлах, весь дизайн-код; писать, что экспортировано из
AuditLens».

Здесь одна точка правды для всех выгрузок:
  • палитра — те же токены, что в интерфейсе (index.html, :root), переведённые
    из OKLCH в sRGB: Excel и Word понимают только HEX;
  • шрифты — Geist (текст), Source Serif 4 (заголовки), JetBrains Mono
    (служебные подписи), как в интерфейсе и PDF. В Word они ВСТРАИВАЮТСЯ в
    документ (у получателя их обычно нет); Excel встраивать не умеет — поэтому
    шапка листа рисуется картинкой с фирменными шрифтами, а в ячейках стоит
    Geist (где его нет, Excel подставит похожий);
  • знак AuditLens — та же «A» с синим штрихом, что в меню;
  • графики для Word — в палитре и шрифтах интерфейса (в Excel графики
    родные, редактируемые);
  • штамп «Экспортировано из AuditLens» — на каждом листе и каждой странице.
"""
from __future__ import annotations

import io
import re
import uuid
import zipfile
from datetime import datetime
from functools import lru_cache
from pathlib import Path

from ..clock import today_msk

# ── палитра (светлая тема интерфейса; OKLCH → sRGB) ──────────────────────────
C = {
    "paper": "FBFAF7",      # oklch(98.5% .004 80)
    "paper2": "F6F4F1",     # oklch(96.8% .005 80)
    "surface": "FFFFFF",
    "ink": "0D1014",        # oklch(17% .01 260)
    "ink2": "3F4348",       # oklch(38% .01 260)
    "ink3": "656970",       # oklch(52% .012 260)
    "ink4": "9499A0",       # oklch(68% .012 260)
    "hair": "E3E1DE",       # oklch(91% .005 80)
    "hair2": "D3D1CD",      # oklch(86% .005 80)
    "accent": "C92F33",     # oklch(55% .19 25) — фирменный красный
    "accent_ink": "AC011A",  # oklch(47% .19 25)
    "accent_soft": "FBF0F0",
    "sber": "006B30",       # oklch(46% .13 152)
    "sber_soft": "E6F0EA",
    "pos": "2C6D3E",
    "warn": "9D6400",
    "warn_ink": "724D0B",
    "warn_soft": "FBF3E4",
    "neg": "CF4040",
    "select": "374960",
    "legal": "683EB6",      # oklch(48% .18 295) — мошеннические схемы
    "info": "036EAE",
    "info_soft": "DFF1FF",
    "mark": "1F4DFF",       # синий штрих знака «A»
}
# Серия графиков: фирменный красный, затем приглушённые тона интерфейса.
SERIES = ["C92F33", "374960", "006B30", "9D6400", "683EB6", "036EAE", "9499A0"]

SANS = "Geist"
SERIF = "Source Serif 4 SemiBold"
MONO = "JetBrains Mono Medium"

STAMP = "Экспортировано из AuditLens"

_FONTS_DIR = Path(__file__).with_name("export_fonts")
_TTF = {
    "sans": "Geist-Regular.ttf", "sans_md": "Geist-Medium.ttf",
    "sans_sb": "Geist-SemiBold.ttf", "sans_b": "Geist-Bold.ttf",
    "serif": "SourceSerif4-Regular.ttf", "serif_sb": "SourceSerif4-SemiBold.ttf",
    "mono": "JetBrainsMono-Medium.ttf",
}


def font_path(key: str) -> Path:
    return _FONTS_DIR / _TTF[key]


def stamp_line(when: datetime | None = None) -> str:
    """«Экспортировано из AuditLens · 26.09.2026 13:05 МСК»."""
    w = when or today_msk()
    return f"{STAMP} · {w:%d.%m.%Y %H:%M} МСК"


def rgb(hex_: str) -> tuple[int, int, int]:
    return tuple(int(hex_[i:i + 2], 16) for i in (0, 2, 4))  # type: ignore[return-value]


# ── картинки: знак, шапка, графики (Pillow + фирменные шрифты) ───────────────

@lru_cache(maxsize=64)
def _pil_font(key: str, size: int):
    from PIL import ImageFont
    return ImageFont.truetype(str(font_path(key)), size)


def _draw_mark(draw, x: float, y: float, size: float, ink: str = C["ink"]) -> None:
    """Знак AuditLens: «A» из меню (viewBox 100×100) — синий штрих и тёмная буква."""
    s = size / 100.0

    def pts(seq):
        return [(x + px * s, y + py * s) for px, py in seq]
    draw.polygon(pts([(47.5, 13), (59.5, 13), (89.5, 89), (75.5, 89)]), fill=rgb(C["mark"]))
    draw.polygon(pts([(47.5, 13), (57.5, 13), (83.5, 89), (66.5, 89), (58.5, 67),
                      (36.5, 67), (27.5, 89), (10.5, 89)]), fill=rgb(ink))
    # «дыра» буквы (fill-rule evenodd в SVG) — цветом фона
    draw.polygon(pts([(47.5, 36), (56.5, 58), (38.5, 58)]), fill=(255, 255, 255))


def _png(img) -> bytes:
    buf = io.BytesIO()
    img.save(buf, format="PNG", optimize=True)
    return buf.getvalue()


def logo_png(height: int = 64) -> bytes:
    """Знак и название — как в меню интерфейса."""
    from PIL import Image, ImageDraw
    k = 3                                   # рисуем крупнее и уменьшаем — ровные края
    h = height * k
    font = _pil_font("sans_sb", int(h * 0.46))
    w = int(h * 1.05 + font.getlength("AuditLens") + h * 0.2)
    img = Image.new("RGB", (w, h), (255, 255, 255))
    d = ImageDraw.Draw(img)
    _draw_mark(d, 0, h * 0.05, h * 0.9)
    d.text((h * 1.0, h * 0.5), "AuditLens", font=font, fill=rgb(C["ink"]), anchor="lm")
    return _png(img.resize((w // k, height), Image.LANCZOS))


def banner_png(*, eyebrow: str, title: str, subtitle: str = "", width: int = 1400) -> bytes:
    """Шапка листа/документа: знак, «AUDITLENS · BANK AUDIT PLATFORM», рубрика,
    заголовок Source Serif 4 и строка сведений. Картинкой — чтобы фирменные
    шрифты были у получателя, даже если у него их нет."""
    from PIL import Image, ImageDraw
    k = 2
    W = width * k
    pad = 36 * k
    f_brand = _pil_font("mono", 13 * k)
    f_eyebrow = _pil_font("mono", 13 * k)
    f_title = _pil_font("serif_sb", 34 * k)
    f_sub = _pil_font("sans", 15 * k)
    title_lines = _wrap(title, f_title, W - 2 * pad)[:2]
    sub_lines = _wrap(subtitle, f_sub, W - 2 * pad)[:2] if subtitle else []
    H = int(pad + 30 * k + 26 * k + 22 * k + 44 * k * len(title_lines)
            + 24 * k * len(sub_lines) + pad * 0.9)
    img = Image.new("RGB", (W, H), rgb(C["paper"]))
    d = ImageDraw.Draw(img)
    d.rectangle([0, 0, W, 5 * k], fill=rgb(C["accent"]))          # фирменная полоса
    y = pad
    _draw_mark(d, pad, y - 4 * k, 28 * k)
    d.text((pad + 36 * k, y + 10 * k), "AUDITLENS · BANK AUDIT PLATFORM", font=f_brand,
           fill=rgb(C["ink"]), anchor="lm")
    stamp = stamp_line().upper()
    d.text((W - pad, y + 10 * k), stamp, font=f_brand, fill=rgb(C["ink3"]), anchor="rm")
    y += 30 * k
    d.line([pad, y, W - pad, y], fill=rgb(C["ink"]), width=k)
    y += 26 * k
    d.text((pad, y), eyebrow.upper(), font=f_eyebrow, fill=rgb(C["accent"]), anchor="lm")
    y += 22 * k
    for ln in title_lines:
        d.text((pad, y), ln, font=f_title, fill=rgb(C["ink"]), anchor="lt")
        y += 44 * k
    for ln in sub_lines:
        d.text((pad, y + 4 * k), ln, font=f_sub, fill=rgb(C["ink3"]), anchor="lt")
        y += 24 * k
    return _png(img.resize((width, H // k), Image.LANCZOS))


def _wrap(text_: str, font, width: float) -> list[str]:
    words, lines, cur = (text_ or "").split(), [], ""
    for w in words:
        nxt = f"{cur} {w}".strip()
        if font.getlength(nxt) <= width or not cur:
            cur = nxt
        else:
            lines.append(cur)
            cur = w
    if cur:
        lines.append(cur)
    return lines


def _num(v: float, unit: str = "") -> str:
    if abs(v - round(v)) < 1e-9:
        s = f"{int(round(v)):,}".replace(",", " ")
    else:
        s = f"{v:.1f}".replace(".", ",")
    return s + (f" {unit}" if unit else "")


def bars_png(title: str, items: list[tuple[str, float]], *, unit: str = "",
             note: str = "", color: str = C["accent"], width: int = 1100,
             max_items: int = 10) -> bytes:
    """Горизонтальные полосы в стиле интерфейса: подпись слева, полоса цвета
    акцента по волосяной дорожке, число моноширинным справа."""
    from PIL import Image, ImageDraw
    items = [(str(a), float(b)) for a, b in items if b][:max_items]
    k = 2
    W = width * k
    pad = 28 * k
    f_title = _pil_font("serif_sb", 20 * k)
    f_lab = _pil_font("sans", 14 * k)
    f_num = _pil_font("mono", 13 * k)
    f_note = _pil_font("sans", 12 * k)
    row = 34 * k
    head = 48 * k
    H = pad + head + row * max(1, len(items)) + (26 * k if note else 0) + pad
    img = Image.new("RGB", (W, H), (255, 255, 255))
    d = ImageDraw.Draw(img)
    d.text((pad, pad), title, font=f_title, fill=rgb(C["ink"]), anchor="lt")
    y = pad + head
    lab_w = int(W * 0.40)
    num_w = 90 * k
    track_x0, track_x1 = pad + lab_w, W - pad - num_w
    top = max((b for _, b in items), default=1) or 1
    for lab, val in items:
        ln = _wrap(lab, f_lab, lab_w - 12 * k)
        d.text((pad, y + row / 2), ln[0] + ("…" if len(ln) > 1 else ""), font=f_lab,
               fill=rgb(C["ink2"]), anchor="lm")
        d.rounded_rectangle([track_x0, y + row * 0.3, track_x1, y + row * 0.7],
                            radius=3 * k, fill=rgb(C["paper2"]))
        x1 = track_x0 + (track_x1 - track_x0) * (val / top)
        d.rounded_rectangle([track_x0, y + row * 0.3, max(track_x0 + 4 * k, x1), y + row * 0.7],
                            radius=3 * k, fill=rgb(color))
        d.text((W - pad, y + row / 2), _num(val, unit), font=f_num, fill=rgb(C["ink"]),
               anchor="rm")
        y += row
    if not items:
        d.text((pad, y + row / 2), "нет данных", font=f_lab, fill=rgb(C["ink3"]), anchor="lm")
        y += row
    if note:
        d.text((pad, y + 10 * k), note, font=f_note, fill=rgb(C["ink3"]), anchor="lt")
    return _png(img.resize((width, H // k), Image.LANCZOS))


# ── Excel ────────────────────────────────────────────────────────────────────

class XL:
    """Оформление листа в стиле интерфейса. Все цвета и шрифты — отсюда.
    Объекты стилей кэшируются: на выгрузке в несколько тысяч строк создание
    шрифта и рамки на каждую ячейку занимало большую часть времени."""

    @staticmethod
    @lru_cache(maxsize=512)
    def font(size: float = 10, bold: bool = False, color: str = "ink", name: str = SANS,
             italic: bool = False, underline: str | None = None):
        from openpyxl.styles import Font
        return Font(name=name, size=size, bold=bold, italic=italic,
                    color=C.get(color, color), underline=underline)

    @staticmethod
    @lru_cache(maxsize=512)
    def fill(color: str):
        from openpyxl.styles import PatternFill
        return PatternFill("solid", fgColor=C.get(color, color))

    @staticmethod
    @lru_cache(maxsize=512)
    def border(bottom: str | None = "hair", top: str | None = None, left: str | None = None,
               right: str | None = None, weight: str = "thin"):
        from openpyxl.styles import Border, Side

        def side(c):
            return Side(style=weight, color=C.get(c, c)) if c else Side()
        return Border(bottom=side(bottom), top=side(top), left=side(left), right=side(right))

    @staticmethod
    def put(ws, row: int, col: int, v):
        """Запись в ячейку. Тексты отзывов чужие: строку «=…» openpyxl записал бы
        формулой — помечаем её текстом. Апостроф-экран здесь не годится: в xlsx
        он становится частью значения и виден в ячейке («'- пункт»)."""
        c = ws.cell(row=row, column=col, value=v)
        if isinstance(v, str) and v.startswith("="):
            c.data_type = "s"
        return c

    @staticmethod
    def sheet_base(ws, *, widths: list[float] | None = None, landscape: bool = True,
                   title: str = "") -> None:
        ws.sheet_view.showGridLines = False
        ws.sheet_properties.tabColor = C["accent"]
        if widths:
            from openpyxl.utils import get_column_letter
            for i, w in enumerate(widths, 1):
                ws.column_dimensions[get_column_letter(i)].width = w
        ps = ws.page_setup
        ps.orientation = "landscape" if landscape else "portrait"
        ps.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_options.horizontalCentered = True
        ws.page_margins.left = ws.page_margins.right = 0.4
        ws.page_margins.top = ws.page_margins.bottom = 0.6
        # колонтитулы печати: знак инструмента и штамп — на каждой странице
        ws.oddHeader.left.text = "AuditLens"
        ws.oddHeader.left.font = f"{SANS},Bold"
        ws.oddHeader.left.size = 9
        ws.oddHeader.right.text = title[:80]
        ws.oddHeader.right.size = 8
        ws.oddFooter.left.text = stamp_line()
        ws.oddFooter.left.size = 8
        ws.oddFooter.right.text = "стр. &P из &N"
        ws.oddFooter.right.size = 8

    @staticmethod
    def banner(ws, png: bytes, *, anchor: str = "A1", rows: int = 7,
               row_height: float = 20.0, width_px: int | None = None) -> int:
        """Картинка-шапка поверх первых строк; возвращает первую свободную строку."""
        from openpyxl.drawing.image import Image as XLImage
        img = XLImage(io.BytesIO(png))
        if width_px:
            ratio = width_px / img.width
            img.width, img.height = width_px, int(img.height * ratio)
        ws.add_image(img, anchor)
        need = int(img.height / (row_height * 96 / 72)) + 1
        for r in range(1, max(rows, need) + 1):
            ws.row_dimensions[r].height = row_height
        return max(rows, need) + 1

    @staticmethod
    def stamp(ws, row: int, col: int = 1, text_: str | None = None, span: int = 8) -> None:
        c = ws.cell(row=row, column=col, value=text_ or stamp_line())
        c.font = XL.font(8.5, color="ink3", name=MONO)
        if span > 1:
            ws.merge_cells(start_row=row, start_column=col, end_row=row,
                           end_column=col + span - 1)

    @staticmethod
    def section(ws, row: int, text_: str, col: int = 1, span: int = 8, eyebrow: str = "") -> int:
        """Заголовок блока: рубрика моноширинным и заголовок Source Serif."""
        if eyebrow:
            e = ws.cell(row=row, column=col, value=eyebrow.upper())
            e.font = XL.font(8.5, color="accent", name=MONO)
            row += 1
        c = ws.cell(row=row, column=col, value=text_)
        c.font = XL.font(15, color="ink", name=SERIF)
        ws.row_dimensions[row].height = 24
        for j in range(col, col + span):
            ws.cell(row=row, column=j).border = XL.border(bottom="ink")
        return row + 2

    @staticmethod
    def kpis(ws, row: int, items: list[tuple[str, object, str]], col: int = 1,
             span_each: int = 2) -> int:
        """Плитки показателей как в интерфейсе: подпись моноширинным капсом,
        крупное число, пояснение. items: (подпись, значение, пояснение)."""
        from openpyxl.styles import Alignment
        for i, (label, value, hint) in enumerate(items):
            c0 = col + i * span_each
            c1 = c0 + span_each - 1
            for r in (row, row + 1, row + 2):
                ws.merge_cells(start_row=r, start_column=c0, end_row=r, end_column=c1)
                for j in range(c0, c1 + 1):
                    cell = ws.cell(row=r, column=j)
                    cell.fill = XL.fill("surface")
                    cell.border = XL.border(bottom="hair2" if r == row + 2 else None,
                                            top="hair2" if r == row else None,
                                            left="hair2" if j == c0 else None,
                                            right="hair2" if j == c1 else None)
            a = ws.cell(row=row, column=c0, value=label.upper())
            a.font = XL.font(8, color="ink3", name=MONO)
            a.alignment = Alignment(indent=1, vertical="bottom")
            v = ws.cell(row=row + 1, column=c0, value=value)
            long_ = len(f"{value:,}" if isinstance(value, (int, float)) else str(value)) > 7
            v.font = XL.font(15 if long_ else 22, bold=True, color="accent" if i == 0 else "ink")
            v.alignment = Alignment(indent=1, vertical="center", horizontal="left")
            if isinstance(value, (int, float)):
                v.number_format = '#,##0' if float(value) == int(value) else '0.0'
            h = ws.cell(row=row + 2, column=c0, value=hint)
            h.font = XL.font(8.5, color="ink3")
            h.alignment = Alignment(indent=1, vertical="top", wrap_text=True)
        ws.row_dimensions[row].height = 18
        ws.row_dimensions[row + 1].height = 34
        ws.row_dimensions[row + 2].height = 28
        return row + 4

    @staticmethod
    def table(ws, row: int, cols: list[str], data: list[list], *, col: int = 1,
              wrap: set[str] | None = None, formats: dict[str, str] | None = None,
              links: set[str] | None = None, name: str | None = None) -> int:
        """Таблица: тёмная шапка (как активная вкладка), зебра цвета бумаги,
        волосяные линии, фильтр. Возвращает строку после таблицы."""
        from openpyxl.styles import Alignment
        from openpyxl.utils import get_column_letter
        wrap = wrap or set()
        formats = formats or {}
        links = links or set()
        for j, k in enumerate(cols, col):
            c = ws.cell(row=row, column=j, value=k)
            c.font = XL.font(9, bold=True, color="surface")
            c.fill = XL.fill("ink")
            c.alignment = Alignment(vertical="center", wrap_text=True, indent=1)
            c.border = XL.border(bottom="ink")
        ws.row_dimensions[row].height = 30
        top = Alignment(vertical="top", wrap_text=False, indent=1)
        topw = Alignment(vertical="top", wrap_text=True, indent=1)
        for i, vals in enumerate(data, row + 1):
            zebra = (i - row) % 2 == 0
            for j, (k, v) in enumerate(zip(cols, vals), col):
                c = XL.put(ws, i, j, v)
                c.font = XL.font(9.5, color="ink")
                c.alignment = topw if k in wrap else top
                c.border = XL.border(bottom="hair")
                if zebra:
                    c.fill = XL.fill("paper")
                if k in formats and v not in (None, ""):
                    c.number_format = formats[k]
                if k in links and isinstance(v, str) and v.startswith("http"):
                    c.hyperlink = v
                    c.font = XL.font(9.5, color="select", underline="single")
        last = row + max(1, len(data))
        ws.auto_filter.ref = (f"{get_column_letter(col)}{row}:"
                              f"{get_column_letter(col + len(cols) - 1)}{last}")
        return last + 1

    @staticmethod
    def style_chart(chart, *, title: str, colors: list[str] | None = None,
                    horizontal: bool = False) -> None:
        """Родной график Excel в палитре интерфейса: без рамки, волосяная сетка,
        заголовок и подписи Geist, серия фирменного красного."""
        from openpyxl.chart.shapes import GraphicalProperties
        from openpyxl.chart.text import RichText
        from openpyxl.drawing.line import LineProperties
        from openpyxl.drawing.text import CharacterProperties, Paragraph, ParagraphProperties
        from openpyxl.drawing.text import Font as DFont

        def txt(size: int, color: str, bold: bool = False):
            cp = CharacterProperties(sz=size * 100, b=bold, solidFill=C.get(color, color),
                                     latin=DFont(typeface=SANS))
            return RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=cp), endParaRPr=cp)])
        chart.title = title
        chart.title.tx.rich.p[0].pPr = ParagraphProperties(defRPr=CharacterProperties(
            sz=1200, b=True, solidFill=C["ink"], latin=DFont(typeface=SANS)))
        chart.style = 2
        chart.legend = None
        chart.graphical_properties = GraphicalProperties(ln=LineProperties(noFill=True))
        colors = colors or SERIES
        for i, s in enumerate(chart.series):
            col = colors[i % len(colors)]
            s.graphicalProperties = GraphicalProperties(
                solidFill=col, ln=LineProperties(solidFill=col, w=28575 if not horizontal
                                                  and chart.tagname == "lineChart" else None))
        for ax in (chart.x_axis, chart.y_axis):
            ax.txPr = txt(9, "ink2")
            ax.graphicalProperties = GraphicalProperties(ln=LineProperties(solidFill=C["hair2"]))
            ax.delete = False
        # сетка — у оси значений (в openpyxl это y и для горизонтальных полос)
        from openpyxl.chart.axis import ChartLines
        chart.y_axis.majorGridlines = ChartLines(spPr=GraphicalProperties(
            ln=LineProperties(solidFill=C["hair"])))
        chart.x_axis.majorGridlines = None


# ── Word ─────────────────────────────────────────────────────────────────────

class WD:
    """Документ Word в стиле интерфейса и PDF-выгрузки отчёта."""

    @staticmethod
    def new(title: str):
        from docx import Document
        from docx.enum.section import WD_ORIENT  # noqa: F401 — книжная по умолчанию
        from docx.shared import Cm, Pt, RGBColor
        doc = Document()
        sec = doc.sections[0]
        sec.page_width, sec.page_height = Cm(21), Cm(29.7)
        sec.left_margin = sec.right_margin = Cm(2)
        sec.top_margin, sec.bottom_margin = Cm(2), Cm(1.8)
        st = doc.styles
        normal = st["Normal"]
        WD._font(normal, SANS, 10.5, C["ink"])
        normal.paragraph_format.space_after = Pt(5)
        normal.paragraph_format.line_spacing = 1.25
        for lvl, size, before in ((0, 26, 0), (1, 17, 16), (2, 13, 12), (3, 11, 10)):
            name = "Title" if lvl == 0 else f"Heading {lvl}"
            h = st[name]
            WD._font(h, SERIF, size, C["ink"], bold=False)
            h.paragraph_format.space_before = Pt(before)
            h.paragraph_format.space_after = Pt(6)
            h.paragraph_format.keep_with_next = True
        for name in ("List Bullet", "List Number"):
            WD._font(st[name], SANS, 10.5, C["ink"])
        doc.core_properties.author = "AuditLens"
        doc.core_properties.last_modified_by = "AuditLens"
        doc.core_properties.title = title
        doc.core_properties.comments = stamp_line()
        doc.core_properties.keywords = "AuditLens"
        _ = RGBColor  # импорт нужен ниже через WD.run
        return doc

    @staticmethod
    def _font(style, name: str, size: float, color: str, bold: bool | None = None) -> None:
        from docx.oxml.ns import qn
        from docx.shared import Pt, RGBColor
        f = style.font
        f.name = name
        f.size = Pt(size)
        f.color.rgb = RGBColor.from_string(color)
        if bold is not None:
            f.bold = bold
        rpr = style.element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = rpr.makeelement(qn("w:rFonts"), {})
            rpr.append(rfonts)
        for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rfonts.set(qn(attr), name)
        # у заголовков Word по умолчанию тема шрифта — снимаем, иначе она победит
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
            if rfonts.get(qn(attr)) is not None:
                del rfonts.attrib[qn(attr)]

    @staticmethod
    def run(par, text_: str, *, font: str = SANS, size: float | None = None,
            color: str | None = None, bold: bool = False, italic: bool = False,
            caps: bool = False, spacing: float | None = None):
        from docx.oxml.ns import qn
        from docx.shared import Pt, RGBColor
        r = par.add_run(text_)
        r.font.name = font
        r._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), font)
        r._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:cs"), font)
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = RGBColor.from_string(C.get(color, color))
        r.bold = bold or None
        r.italic = italic or None
        r.font.all_caps = caps or None
        if spacing is not None:
            rpr = r._element.get_or_add_rPr()
            sp = rpr.makeelement(qn("w:spacing"), {qn("w:val"): str(int(spacing * 20))})
            rpr.append(sp)
        return r

    @staticmethod
    def eyebrow(doc_or_cell, text_: str, color: str = "accent"):
        p = doc_or_cell.add_paragraph()
        WD.run(p, text_.upper(), font=MONO, size=7.5, color=color, spacing=0.6)
        p.paragraph_format.space_after = WD._pt(2)
        return p

    @staticmethod
    def _pt(v: float):
        from docx.shared import Pt
        return Pt(v)

    @staticmethod
    def shade(cell, color: str) -> None:
        from docx.oxml.ns import qn
        tcpr = cell._element.get_or_add_tcPr()
        shd = tcpr.makeelement(qn("w:shd"), {qn("w:val"): "clear", qn("w:color"): "auto",
                                             qn("w:fill"): C.get(color, color)})
        tcpr.append(shd)

    @staticmethod
    def cell_borders(cell, **sides) -> None:
        """sides: top/left/bottom/right = (цвет, толщина в восьмых пункта) или None."""
        from docx.oxml.ns import qn
        tcpr = cell._element.get_or_add_tcPr()
        b = tcpr.makeelement(qn("w:tcBorders"), {})
        for side in ("top", "left", "bottom", "right"):
            v = sides.get(side)
            el = b.makeelement(qn(f"w:{side}"), (
                {qn("w:val"): "single", qn("w:sz"): str(v[1]), qn("w:space"): "0",
                 qn("w:color"): C.get(v[0], v[0])} if v else {qn("w:val"): "nil"}))
            b.append(el)
        tcpr.append(b)

    @staticmethod
    def rule(doc, color: str = "ink", size: int = 8) -> None:
        """Горизонтальная линия под абзацем (как волосяные линии интерфейса)."""
        from docx.oxml.ns import qn
        p = doc.add_paragraph()
        ppr = p._element.get_or_add_pPr()
        bdr = ppr.makeelement(qn("w:pBdr"), {})
        bdr.append(bdr.makeelement(qn("w:bottom"), {qn("w:val"): "single",
                                                    qn("w:sz"): str(size), qn("w:space"): "1",
                                                    qn("w:color"): C.get(color, color)}))
        ppr.append(bdr)
        p.paragraph_format.space_after = WD._pt(6)

    @staticmethod
    def header_footer(doc, title: str) -> None:
        """Колонтитулы: «AUDITLENS · …» сверху, штамп и «стр. N из M» снизу."""
        from docx.oxml.ns import qn
        sec = doc.sections[0]
        sec.different_first_page_header_footer = True     # на обложке — без колонтитулов
        hp = sec.header.paragraphs[0]
        WD.run(hp, "AUDITLENS", font=MONO, size=7.5, color="ink", spacing=0.8)
        WD.run(hp, f"  ·  {title[:90]}", font=MONO, size=7.5, color="ink3")
        # Одна строка без табуляции: стандартный стиль колонтитула несёт свои
        # позиции табуляции, и «стр.» у LibreOffice прилипал к штампу.
        fp = sec.footer.paragraphs[0]
        WD.run(fp, stamp_line() + "  ·  стр. ", font=MONO, size=7.5, color="ink3")
        for part in ("PAGE", " из ", "NUMPAGES"):
            if part.strip() not in ("PAGE", "NUMPAGES"):
                WD.run(fp, part, font=MONO, size=7.5, color="ink3")
                continue
            r = WD.run(fp, "", font=MONO, size=7.5, color="ink3")
            for kind in ("begin", "instr", "separate", "text", "end"):
                if kind == "instr":
                    el = r._element.makeelement(qn("w:instrText"), {})
                    el.set(qn("xml:space"), "preserve")
                    el.text = f" {part} "
                elif kind == "text":
                    el = r._element.makeelement(qn("w:t"), {})
                    el.text = "1"
                else:
                    el = r._element.makeelement(qn("w:fldChar"), {qn("w:fldCharType"): kind})
                r._element.append(el)

    @staticmethod
    def cover(doc, *, eyebrow: str, title: str, meta: list[tuple[str, str]],
              ident: str = "") -> None:
        """Обложка как у PDF-отчёта: линейка, «AUDITLENS · BANK AUDIT PLATFORM»,
        рубрика, заголовок Source Serif, таблица сведений, штамп."""
        from docx.enum.table import WD_TABLE_ALIGNMENT
        from docx.shared import Cm, Pt
        doc.add_paragraph().paragraph_format.space_after = Pt(60)
        WD.rule(doc, "ink", 12)
        t = doc.add_table(rows=1, cols=2)
        t.alignment = WD_TABLE_ALIGNMENT.LEFT
        l, r = t.rows[0].cells
        l.width, r.width = Cm(12), Cm(5)
        lp = l.paragraphs[0]
        lp.add_run().add_picture(io.BytesIO(logo_png(64)), height=Cm(0.8))
        lp2 = l.add_paragraph()
        WD.run(lp2, "AUDITLENS · BANK AUDIT PLATFORM", font=MONO, size=8, color="ink",
               spacing=1.2)
        rp = r.paragraphs[0]
        rp.alignment = 2
        WD.run(rp, ident, font=MONO, size=8, color="ink3")
        WD.rule(doc, "ink", 4)
        doc.add_paragraph().paragraph_format.space_after = Pt(24)
        WD.eyebrow(doc, eyebrow, color="ink3")
        # не стиль «Title»: в шаблоне Word у него синяя линия темы оформления
        tp = doc.add_paragraph()
        WD.run(tp, title, font=SERIF, size=26, color="ink")
        tp.paragraph_format.line_spacing = 1.1
        tp.paragraph_format.space_after = Pt(28)
        mt = doc.add_table(rows=len(meta), cols=2)
        for (k, v), row in zip(meta, mt.rows):
            a, b = row.cells
            a.width, b.width = Cm(4), Cm(13)
            WD.run(a.paragraphs[0], k.upper(), font=MONO, size=7.5, color="ink3", spacing=0.6)
            WD.run(b.paragraphs[0], v, font=SANS, size=10.5, color="ink")
        doc.add_paragraph().paragraph_format.space_after = Pt(18)
        sp = doc.add_paragraph()
        WD.run(sp, stamp_line(), font=MONO, size=8, color="accent")
        doc.add_page_break()

    @staticmethod
    def kpis(doc, items: list[tuple[str, str, str]]) -> None:
        """Плитки показателей: подпись моноширинным, число Source Serif, пояснение."""
        from docx.shared import Cm
        n = len(items)
        t = doc.add_table(rows=1, cols=n)
        for (label, value, hint), cell in zip(items, t.rows[0].cells):
            cell.width = Cm(17 / n)
            WD.shade(cell, "paper")
            WD.cell_borders(cell, top=("accent", 12), left=None, right=None,
                            bottom=("hair", 4))
            p = cell.paragraphs[0]
            WD.run(p, label.upper(), font=MONO, size=7, color="ink3", spacing=0.5)
            p2 = cell.add_paragraph()
            WD.run(p2, value, font=SERIF, size=22, color="ink")
            p2.paragraph_format.space_after = WD._pt(0)
            if hint:
                p3 = cell.add_paragraph()
                WD.run(p3, hint, font=SANS, size=8, color="ink3")
        doc.add_paragraph()

    @staticmethod
    def picture(doc, png: bytes, width_cm: float = 17.0) -> None:
        from docx.shared import Cm
        doc.add_paragraph().add_run().add_picture(io.BytesIO(png), width=Cm(width_cm))


# ── встраивание шрифтов в .docx ──────────────────────────────────────────────
# Word показывает встроенные TrueType-шрифты, даже если у получателя их нет.
# Формат (ECMA-376, ч. 1, §17.8.1): файл шрифта «обфусцирован» — первые 32 байта
# XOR с ключом-GUID (байты GUID в обратном порядке), лежит в word/fonts/*.odttf,
# связан из word/fontTable.xml (w:embedRegular / w:embedBold с w:fontKey).

_EMBED = [  # имя шрифта в документе → файлы начертаний
    (SANS, {"embedRegular": "sans", "embedBold": "sans_b"}),
    (SERIF, {"embedRegular": "serif_sb"}),
    (MONO, {"embedRegular": "mono"}),
]
_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_FONT_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/font"


def obfuscate(data: bytes, key: str) -> bytes:
    g = bytes.fromhex(key.strip("{}").replace("-", ""))[::-1]
    b = bytearray(data)
    for i in range(32):
        b[i] ^= g[i % 16]
    return bytes(b)


def embed_fonts(docx: bytes) -> bytes:
    """Встраивает фирменные шрифты в собранный документ Word."""
    src = zipfile.ZipFile(io.BytesIO(docx))
    files = {n: src.read(n) for n in src.namelist()}
    ft = files["word/fontTable.xml"].decode("utf-8")
    rels_name = "word/_rels/fontTable.xml.rels"
    rels = files.get(rels_name, b"").decode("utf-8") or (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '</Relationships>')
    if 'xmlns:r=' not in ft.split(">", 2)[1]:
        ft = re.sub(r"(<w:fonts\b)", rf'\1 xmlns:r="{_R}"', ft, count=1)
    n = 0
    new_rels = []
    for name, faces in _EMBED:
        embeds = []
        for tag, key in faces.items():
            n += 1
            guid = "{" + str(uuid.uuid4()).upper() + "}"
            rid = f"rIdAlFont{n}"
            files[f"word/fonts/font{n}.odttf"] = obfuscate(font_path(key).read_bytes(), guid)
            new_rels.append(f'<Relationship Id="{rid}" Type="{_FONT_REL}" '
                            f'Target="fonts/font{n}.odttf"/>')
            embeds.append(f'<w:{tag} r:id="{rid}" w:fontKey="{guid}"/>')
        block = (f'<w:font w:name="{name}"><w:charset w:val="CC"/>'
                 f'<w:family w:val="auto"/><w:pitch w:val="variable"/>{"".join(embeds)}</w:font>')
        ft = re.sub(rf'<w:font w:name="{re.escape(name)}".*?</w:font>', "", ft, flags=re.DOTALL)
        ft = ft.replace("</w:fonts>", block + "</w:fonts>")
    rels = rels.replace("</Relationships>", "".join(new_rels) + "</Relationships>")
    files["word/fontTable.xml"] = ft.encode("utf-8")
    files[rels_name] = rels.encode("utf-8")
    ct = files["[Content_Types].xml"].decode("utf-8")
    if 'Extension="odttf"' not in ct:
        ct = ct.replace("<Types ", "<Types ", 1).replace(
            "</Types>", '<Default Extension="odttf" '
            'ContentType="application/vnd.openxmlformats-officedocument.obfuscatedFont"/>'
            "</Types>")
    files["[Content_Types].xml"] = ct.encode("utf-8")
    st = files["word/settings.xml"].decode("utf-8")
    if "embedTrueTypeFonts" not in st:
        st = re.sub(r"(<w:settings\b[^>]*>)", r"\1<w:embedTrueTypeFonts/>", st, count=1)
    files["word/settings.xml"] = st.encode("utf-8")
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        # [Content_Types].xml — первым: так ждут строгие читатели пакета
        z.writestr("[Content_Types].xml", files.pop("[Content_Types].xml"))
        for name_, data in files.items():
            z.writestr(name_, data)
    return out.getvalue()


def docx_bytes(doc) -> bytes:
    buf = io.BytesIO()
    doc.save(buf)
    return embed_fonts(buf.getvalue())
