"""Аудит-дело, этап 5: команды и полная выгрузка.

Пересчёт доступа «через команду» (SQL) проверяется на живой базе.
"""
from __future__ import annotations

import io
import zipfile
from datetime import datetime, timezone

import openpyxl

from bank_audit.web import case_export as X
from bank_audit.web import notices as N
from bank_audit.web import userdata as U

CASE = {"case_id": 1, "title": "Тест", "owner_name": "Владелец Один", "status_label": "В работе",
        "members": [{"username": "o", "name": "Владелец Один", "role": "owner", "role_label": "владелец"},
                    {"username": "e", "name": "Елена Волкова", "role": "editor",
                     "role_label": "может добавлять", "team_id": 3, "team_name": "Карты",
                     "added_at": datetime(2026, 10, 2, tzinfo=timezone.utc)},
                    {"username": "v", "name": "Ольга Смирнова", "role": "viewer",
                     "role_label": "только смотрит", "added_at": datetime(2026, 10, 2, tzinfo=timezone.utc)}],
        "items": [{"item_id": 11, "kind": "news", "url": "https://x", "title": "Новость",
                   "meta": {"summary": "почему важно"}}]}
TALK = [{"msg_id": 1, "item_id": None, "name": "Елена Волкова", "body": "Начнём с [1]",
         "created_at": datetime(2026, 10, 3, 9, 0, tzinfo=timezone.utc), "reply_to": None},
        {"msg_id": 2, "item_id": 11, "name": "Владелец Один", "body": "Согласна",
         "created_at": datetime(2026, 10, 3, 10, 0, tzinfo=timezone.utc), "reply_to": 1, "edited_at": 1},
        {"msg_id": 3, "item_id": None, "name": "Ольга Смирнова", "body": "", "deleted": True,
         "created_at": None, "reply_to": None}]
HIST = [{"created_at": datetime(2026, 10, 3, 11, tzinfo=timezone.utc), "text": "Подключена команда «Карты»",
         "who": "Владелец О."},
        {"created_at": datetime(2026, 10, 1, 9, tzinfo=timezone.utc), "text": "Дело создано", "who": "Владелец О."}]


def test_people_rows_show_how_added_without_logins():
    rows = X.people_rows(CASE)
    assert rows[0][:3] == ["Владелец Один", "владелец", "создал дело"]
    assert rows[1][2] == "команда «Карты»" and rows[2][2] == "лично"
    assert all("e" != c for r in rows for c in r)


def test_talk_rows_skip_deleted_number_items_and_name_replies():
    rows = X.talk_rows(CASE, TALK)
    assert len(rows) == 2
    assert rows[0] == ["03.10.2026 12:00", "Елена В.", "", "", "Начнём с [1]"]
    assert rows[1] == ["03.10.2026 13:00", "Владелец О.", "[1]", "Елена В.", "Согласна (изменено)"]


def test_history_rows_are_chronological_msk():
    assert [r[1] for r in X.history_rows(HIST)] == ["Дело создано", "Подключена команда «Карты»"]
    assert X.history_rows(HIST)[0][0] == "01.10.2026 12:00"


def test_xlsx_sheets_follow_the_choice():
    full = openpyxl.load_workbook(io.BytesIO(X.to_xlsx(CASE, talk=TALK, history=HIST)))
    assert {"Участники", "Обсуждение", "История"} <= set(full.sheetnames)
    lean = openpyxl.load_workbook(io.BytesIO(X.to_xlsx(CASE)))
    assert "Участники" in lean.sheetnames and "Обсуждение" not in lean.sheetnames
    ws = full["Обсуждение"]
    vals = [c.value for row in ws.iter_rows() for c in row if c.value]
    assert "Согласна (изменено)" in vals


def test_docx_appendices_only_when_chosen():
    def body(**kw):
        return zipfile.ZipFile(io.BytesIO(X.to_docx(CASE, **kw))).read("word/document.xml").decode()
    full, lean = body(talk=TALK, history=HIST), body()
    assert "Обсуждение дела" in full and "История дела" in full and "Начнём с [1]" in full
    assert "Обсуждение дела" not in lean and "История дела" not in lean
    assert "КТО ВЕДЁТ ДЕЛО" in lean.upper() and "команда «Карты»" in lean


def test_team_history_lines():
    n = {"e": "Елена В."}.get
    assert U._ev_text("team_added", {"team": "Карты", "role": "editor"}, n) \
        == "Подключена команда «Карты» — может добавлять"
    assert U._ev_text("team_role", {"team": "Карты", "role": "viewer"}, n) == "Права команды «Карты»: только смотрит"
    assert U._ev_text("team_removed", {"team": "Карты"}, n) == "Отключена команда «Карты»"
    assert U._ev_text("team_removed", {"team": "Карты", "deleted": True}, n) == "Команда «Карты» удалена её создателем"
    assert U._ev_text("member_added", {"member": "e", "role": "editor", "team": "Карты"}, n) \
        == "Новый участник из команды «Карты»: Елена В. — может добавлять"
    assert U._ev_text("member_removed", {"member": "e", "team": True}, n) == "Доступ через команду снят: Елена В."


def test_notice_names_the_team():
    assert N.title_of("case_added", {"case": "К", "team": "Карты"}) == "Вас добавили в дело «К» — команда «Карты»"
    assert N.title_of("case_added", {"case": "К"}) == "Вас добавили в дело «К»"


def test_only_own_team_attaches_to_own_case(monkeypatch):
    monkeypatch.setattr(U, "_owns_case", lambda c, u: True)
    monkeypatch.setattr(U, "_own_team", lambda t, u: None)          # чужая команда
    assert U.attach_team(1, "o", 9, "editor") == "подключить можно только свою команду"
    monkeypatch.setattr(U, "_owns_case", lambda c, u: False)
    assert U.attach_team(1, "x", 9, "editor") == "управлять доступом может только владелец дела"
    assert U.attach_team(1, "o", 9, "admin") == "неизвестная роль"


def test_team_needs_a_name_and_owner_check(monkeypatch):
    try:
        U.create_team("o", "   ")
        raise AssertionError("без названия нельзя")
    except U.CaseError:
        pass
    monkeypatch.setattr(U, "_own_team", lambda t, u: None)
    for fn in (lambda: U.update_team("x", 1, name="Новое"), lambda: U.delete_team("x", 1)):
        try:
            fn()
            raise AssertionError("чужую команду не меняют")
        except U.CaseError:
            pass


def test_clean_members_drops_owner_unknown_and_dupes(monkeypatch):
    monkeypatch.setattr(U, "_rows", lambda sql, p=None: [{"username": "a"}, {"username": "b"}])
    assert U._clean_members("o", ["a", "o", "a", "zz", "b", ""]) == ["a", "b"]
